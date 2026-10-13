from __future__ import annotations

import csv
import codecs
import json
import math
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

from backend.app.services.query import connect


TABULAR_SUFFIXES = {".csv", ".xlsx", ".xls"}
REQUIRED_FIELDS = ("mmsi", "timestamp", "lon", "lat", "speed", "course")
OPTIONAL_FIELDS = ("ship_type", "length", "width", "draught", "heading")
STANDARD_FIELDS = (*REQUIRED_FIELDS, *OPTIONAL_FIELDS)
FIELD_LABELS = {
    "mmsi": "MMSI", "timestamp": "时间", "lon": "经度", "lat": "纬度",
    "speed": "航速", "course": "对地航向", "ship_type": "船舶类型",
    "length": "船长", "width": "船宽", "draught": "吃水", "heading": "船首向",
}
# Header matching ignores case, whitespace, underscores and punctuation.
# COG/course and true heading are deliberately separate fields.
FIELD_ALIASES = {
    "mmsi": ("mmsi", "mmsi_no", "mmsi_number", "vessel_mmsi", "船舶编号", "船舶识别码", "海上移动业务识别码"),
    "timestamp": ("timestamp", "updatetime", "update_time", "datetime", "date_time", "time", "record_time", "position_time", "utc_time", "时间", "更新时间", "采集时间", "定位时间", "时间戳"),
    "lon": ("lon", "longitude", "lng", "long", "经度"),
    "lat": ("lat", "latitude", "纬度"),
    "speed": ("speed", "sog", "speed_over_ground", "vessel_speed", "speed_knots", "航速", "航速节", "对地速度"),
    "course": ("course", "cog", "course_over_ground", "航向", "对地航向", "航向角"),
    "ship_type": ("ship_type", "shiptype", "type", "vessel_type", "vesseltype", "船舶类型", "船型", "类型"),
    "length": ("length", "ship_length", "vessel_length", "loa", "船长", "总长", "船长米"),
    "width": ("width", "ship_width", "vessel_width", "beam", "船宽", "船宽米"),
    "draught": ("draught", "draft", "ship_draught", "ship_draft", "吃水", "吃水深度", "吃水米"),
    "heading": ("heading", "true_heading", "trueheading", "艏向", "船首向", "船艏向"),
}
NULL_TEXT = {"", "null", "none", "nan", "nat", "n/a", "na", "-", "--", "未知", "无"}


def validate_options(source_timezone: str = "Asia/Shanghai", csv_encoding: str = "auto") -> None:
    try:
        ZoneInfo(source_timezone)
    except (ZoneInfoNotFoundError, ValueError) as exc:
        raise ValueError(f"无法识别时间时区：{source_timezone}") from exc
    if csv_encoding not in {"auto", "utf-8-sig", "utf-8", "gb18030"}:
        raise ValueError("CSV 编码应为 auto、utf-8-sig、utf-8 或 gb18030。")


def normalize_header(value: Any) -> str:
    return "".join(char for char in str(value).strip().lstrip("\ufeff").casefold() if char.isalnum())


def resolve_fields(columns: list[str]) -> dict[str, str]:
    normalized = [normalize_header(column) for column in columns]
    duplicate = [key for key, count in Counter(normalized).items() if key and count > 1]
    if duplicate:
        raise ValueError(f"存在重复或归一化后重名的表头：{', '.join(duplicate)}")
    mapping = {}
    for field, aliases in FIELD_ALIASES.items():
        accepted = {normalize_header(alias) for alias in aliases}
        matches = [column for column, key in zip(columns, normalized) if key in accepted]
        if len(matches) > 1:
            raise ValueError(f"{FIELD_LABELS[field]}匹配到多个字段：{', '.join(matches)}。请保留一个明确的字段名。")
        if matches:
            mapping[field] = matches[0]
    missing = [f"{FIELD_LABELS[field]} ({field})" for field in REQUIRED_FIELDS if field not in mapping]
    if missing:
        raise ValueError("缺少必需字段：" + "、".join(missing))
    return mapping


def _headers(values: tuple[Any, ...] | list[Any]) -> list[str]:
    columns = [str(value).strip() if value is not None else "" for value in values]
    while columns and not columns[-1]:
        columns.pop()
    return [column or f"__unnamed_{index + 1}" for index, column in enumerate(columns)]


def _csv_options(path: Path, encoding: str) -> tuple[str, str, list[str]]:
    with path.open("rb") as handle:
        sample = handle.read(65536)
    if encoding == "auto":
        encoding = "utf-8-sig"
        try:
            codecs.getincrementaldecoder(encoding)().decode(sample, final=False)
        except UnicodeDecodeError:
            encoding = "gb18030"
    try:
        text = codecs.getincrementaldecoder(encoding)().decode(sample, final=False)
    except UnicodeDecodeError as exc:
        raise ValueError("CSV 编码不匹配，请选择 UTF-8 或 GB18030 后重试。") from exc
    # A sample may end mid-character; use a full first line for the header.
    try:
        delimiter = csv.Sniffer().sniff(text, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    with path.open("r", encoding=encoding, newline="") as handle:
        headers = next(csv.reader(handle, delimiter=delimiter), [])
    return encoding, delimiter, _headers(headers)


def source_chunks(
    path: Path, csv_encoding: str = "auto", sheet_name: str | None = None, chunk_rows: int = 100_000,
) -> Iterator[tuple[str, list[str], pd.DataFrame, dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        encoding, delimiter, columns = _csv_options(path, csv_encoding)
        resolve_fields(columns)
        reader = pd.read_csv(path, encoding=encoding, sep=delimiter, dtype="string", keep_default_na=False, chunksize=chunk_rows)
        with reader:
            for frame in reader:
                frame = frame.iloc[:, :len(columns)]
                frame.columns = columns
                if len(frame):
                    yield "CSV", columns, frame, {"encoding": encoding, "delimiter": delimiter}
        return
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name and sheet_name not in workbook.sheetnames:
                raise ValueError(f"找不到工作表：{sheet_name}")
            for sheet in workbook.worksheets:
                if sheet_name and sheet.title != sheet_name:
                    continue
                rows = sheet.iter_rows(values_only=True)
                header = next((row for row in rows if any(value is not None for value in row)), None)
                if header is None:
                    continue
                columns = _headers(header)
                # Ignore genuinely empty worksheets, including the sample's Sheet1.
                first = next((row for row in rows if any(value is not None for value in row)), None)
                if first is None:
                    continue
                try:
                    resolve_fields(columns)
                except ValueError as exc:
                    raise ValueError(f"工作表“{sheet.title}”：{exc}") from exc
                batch = [first[:len(columns)]]
                for row in rows:
                    if not any(value is not None for value in row):
                        continue
                    batch.append(row[:len(columns)])
                    if len(batch) >= chunk_rows:
                        yield sheet.title, columns, pd.DataFrame(batch, columns=columns), {}
                        batch = []
                if batch:
                    yield sheet.title, columns, pd.DataFrame(batch, columns=columns), {}
        finally:
            workbook.close()
        return
    if suffix == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(path, on_demand=True)
        try:
            if sheet_name and sheet_name not in workbook.sheet_names():
                raise ValueError(f"找不到工作表：{sheet_name}")
            for sheet in workbook.sheets():
                if sheet_name and sheet.name != sheet_name or sheet.nrows < 2:
                    continue
                columns = _headers(sheet.row_values(0))
                try:
                    resolve_fields(columns)
                except ValueError as exc:
                    raise ValueError(f"工作表“{sheet.name}”：{exc}") from exc
                for start in range(1, sheet.nrows, chunk_rows):
                    batch = []
                    for index in range(start, min(start + chunk_rows, sheet.nrows)):
                        row = sheet.row_values(index)[:len(columns)]
                        for col, cell in enumerate(sheet.row(index)[:len(columns)]):
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                row[col] = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                        if any(value not in (None, "") for value in row):
                            batch.append(row)
                    if batch:
                        yield sheet.name, columns, pd.DataFrame(batch, columns=columns), {}
        finally:
            workbook.release_resources()
        return
    raise ValueError("当前支持 CSV、Excel (.xlsx/.xls) 和原有 PKL/PICKLE 文件。")


def _missing(series: pd.Series) -> pd.Series:
    return series.isna() | series.astype("string").str.strip().str.casefold().isin(NULL_TEXT)


def parse_timestamps(values: pd.Series, source_timezone: str) -> pd.Series:
    result = pd.Series(0, index=values.index, dtype="int64")
    text = values.astype("string").str.strip()
    numbers = pd.to_numeric(text, errors="coerce")
    finite = numbers.notna() & np.isfinite(numbers.to_numpy(dtype=float, na_value=np.nan))

    def assign(parsed: pd.Series) -> None:
        good = parsed.notna()
        if good.any():
            result.loc[parsed.index[good]] = (parsed.loc[good].dt.as_unit("ns").astype("int64") // 1_000_000_000).to_numpy()

    for unit, low, high in (("s", 0, 1e11), ("ms", 1e11, 1e14), ("us", 1e14, 1e17), ("ns", 1e17, math.inf)):
        selected = finite & numbers.abs().ge(low) & numbers.abs().lt(high)
        if selected.any():
            assign(pd.to_datetime(numbers.loc[selected], unit=unit, utc=True, errors="coerce"))
    remaining = ~finite & ~_missing(values)
    aware = remaining & text.str.contains(r"(?:Z|[+-]\d{2}:?\d{2})$", regex=True, case=False, na=False)
    if aware.any():
        assign(pd.to_datetime(text.loc[aware], format="mixed", utc=True, errors="coerce"))
    naive = remaining & ~aware
    if naive.any():
        parsed = pd.to_datetime(text.loc[naive], format="mixed", errors="coerce")
        assign(parsed.dt.tz_localize(source_timezone, ambiguous="NaT", nonexistent="NaT").dt.tz_convert("UTC"))
    return result


def normalize_frame(frame: pd.DataFrame, mapping: dict[str, str], source_timezone: str) -> tuple[pd.DataFrame, dict[str, int]]:
    normalized = pd.DataFrame(index=frame.index)
    issues = {}
    for field in STANDARD_FIELDS:
        if field not in mapping:
            normalized[field] = -1 if field == "ship_type" else np.nan
            continue
        values = frame[mapping[field]]
        if field == "timestamp":
            converted = parse_timestamps(values, source_timezone)
            invalid = converted.eq(0) & ~_missing(values)
        else:
            converted = pd.to_numeric(values.mask(_missing(values)), errors="coerce").astype(float)
            invalid = (~np.isfinite(converted)) & ~_missing(values)
            converted = converted.where(np.isfinite(converted))
            if field in {"mmsi", "ship_type"}:
                # Reject fractional identifiers rather than silently truncating them.
                bad_integer = converted.notna() & converted.mod(1).ne(0)
                if field == "ship_type":
                    bad_integer |= converted.notna() & ~converted.between(-32768, 32767)
                else:
                    bad_integer |= converted.notna() & ~converted.between(0, 999999999)
                invalid |= bad_integer
                converted = converted.mask(bad_integer).fillna(-1 if field == "ship_type" else 0).astype("int64")
        normalized[field] = converted
        issues[field] = int(invalid.sum())
    return normalized, issues


def inspect_tabular(path: Path, source_timezone: str = "Asia/Shanghai", csv_encoding: str = "auto", sheet_name: str | None = None) -> dict[str, Any]:
    validate_options(source_timezone, csv_encoding)
    chunks = source_chunks(path, csv_encoding, sheet_name, chunk_rows=5)
    try:
        sheet, columns, frame, options = next(chunks)
    except StopIteration as exc:
        raise ValueError("文件没有可导入的数据行。") from exc
    finally:
        chunks.close()
    mapping = resolve_fields(columns)
    normalized, issues = normalize_frame(frame.head(5), mapping, source_timezone)
    warnings = [f"未提供{FIELD_LABELS[field]}，按缺失值保留。" for field in OPTIONAL_FIELDS if field not in mapping]
    if "ship_type" in mapping and normalized["ship_type"].eq(-1).all():
        warnings.append("预览行的船舶类型为空或不可识别，将标记为未知类型 (-1)。")
    return {
        "source": str(path), "format": path.suffix.lower().lstrip("."), "sheet": sheet,
        "source_timezone": source_timezone, "columns": columns, "mapping": mapping,
        "field_labels": FIELD_LABELS, "required_fields": list(REQUIRED_FIELDS), "warnings": warnings,
        "ignored_columns": [column for column in columns if column not in mapping.values()],
        "preview": json.loads(normalized.to_json(orient="records")), "preview_invalid_values": issues, **options,
    }


def build_tabular_parquet(path: Path, output_path: Path, manifest_path: Path, dataset_id: str, dataset_name: str,
                          source_timezone: str = "Asia/Shanghai", csv_encoding: str = "auto", sheet_name: str | None = None,
                          progress=None) -> dict[str, Any]:
    validate_options(source_timezone, csv_encoding)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(output_path.name + ".import.tmp")
    sorted_path = output_path.with_name(output_path.name + ".sorted.tmp")
    started = time.perf_counter()
    schema = pa.schema([
        pa.field("mmsi", pa.int64()), pa.field("timestamp", pa.int64()),
        pa.field("lon", pa.float64()), pa.field("lat", pa.float64()),
        pa.field("course", pa.float32()), pa.field("speed", pa.float32()),
        pa.field("ship_type", pa.int16()), pa.field("length", pa.float32()), pa.field("width", pa.float32()),
        pa.field("draught", pa.float32()), pa.field("heading", pa.float32()), pa.field("trajectory_id", pa.int32()),
    ])
    writer = None
    vessels: dict[int, int] = {}
    mappings = {}
    invalid = Counter()
    row_count = 0
    try:
        for sheet, columns, frame, options in source_chunks(path, csv_encoding, sheet_name):
            mapping = resolve_fields(columns)
            mappings[sheet] = {"mapping": mapping, "columns": columns, **options}
            normalized, issues = normalize_frame(frame, mapping, source_timezone)
            invalid.update(issues)
            normalized["trajectory_id"] = [vessels.setdefault(int(value), len(vessels)) for value in normalized["mmsi"]]
            table = pa.Table.from_pandas(normalized[schema.names], schema=schema, preserve_index=False)
            if writer is None:
                writer = pq.ParquetWriter(temporary, schema, compression="zstd")
            writer.write_table(table)
            row_count += len(frame)
            if progress:
                progress({"processed_rows": row_count, "sheet": sheet, "elapsed_seconds": round(time.perf_counter() - started, 2)})
        if writer is None:
            raise ValueError("文件没有可导入的数据行。")
        writer.close()
        writer = None
        escaped_source = str(temporary).replace("\\", "/").replace("'", "''")
        escaped_target = str(sorted_path).replace("\\", "/").replace("'", "''")
        with connect() as con:
            con.execute(f"COPY (SELECT * FROM read_parquet('{escaped_source}') ORDER BY mmsi,timestamp) TO '{escaped_target}' (FORMAT PARQUET, COMPRESSION ZSTD)")
            ranges = con.execute(f"SELECT min(lon),max(lon),min(lat),max(lat),min(timestamp),max(timestamp),min(speed),max(speed) FROM read_parquet('{escaped_source}')").fetchone()
        manifest = {
            "dataset_id": dataset_id, "dataset_name": dataset_name, "data_role": "source",
            "source": str(path), "parquet": str(output_path), "source_format": path.suffix.lower().lstrip("."),
            "source_timezone": source_timezone, "field_mappings": mappings, "conversion_invalid_values": dict(invalid),
            "columns": schema.names, "row_count": row_count, "trajectory_count": len(vessels), "vessel_count": len(vessels),
            "created_at": datetime.now(timezone.utc).isoformat(), "elapsed_seconds": round(time.perf_counter() - started, 2),
            "bounds": dict(zip(("lon_min", "lon_max", "lat_min", "lat_max"), ranges[:4])),
            "time_range": dict(zip(("timestamp_min", "timestamp_max"), ranges[4:6])),
            "speed_range": dict(zip(("speed_min", "speed_max"), ranges[6:8])),
        }
        sorted_path.replace(output_path)
        manifest_path.parent.mkdir(parents=True, exist_ok=True)
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        return manifest
    finally:
        if writer is not None:
            writer.close()
        for scratch in (temporary, sorted_path):
            if scratch.exists():
                scratch.unlink()
