from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from backend.app.services import analysis, directions, exporter, query
from backend.app.services.query import QueryFilters, normalize_spatial_filter


class HistogramTests(unittest.TestCase):
    start = 1_700_000_000

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.parquet = self.root / "points.parquet"
        self.manifest = {}
        for target, name, value in (
            (analysis, "dataset_ready", lambda: True),
            (analysis, "read_active_manifest", lambda: self.manifest),
            (query, "active_parquet_path", lambda: self.parquet),
            (query, "read_active_manifest", lambda: self.manifest),
            (directions, "read_active_manifest", lambda: self.manifest),
            (exporter, "EXPORT_DIR", self.root / "exports"),
        ):
            patcher = patch.object(target, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def point(self, offset=0, **extra):
        row = {"mmsi": 413000001, "trajectory_id": 1,
               "timestamp": self.start + offset, "lon": 110.0, "lat": 20.0,
               "speed": 5.0, "course": 90.0, "heading": 90.0,
               "ship_type": 70, "length": 100.0, "width": 20.0}
        row.update(extra)
        return row

    def write(self, rows):
        schema = pa.schema([
            ("mmsi", pa.int64()), ("trajectory_id", pa.int64()),
            ("timestamp", pa.int64()), ("lon", pa.float64()), ("lat", pa.float64()),
            ("speed", pa.float64()), ("course", pa.float64()), ("heading", pa.float64()),
            ("ship_type", pa.int16()), ("length", pa.float64()), ("width", pa.float64()),
        ])
        pq.write_table(pa.Table.from_pylist(rows, schema=schema), self.parquet)
        self.manifest = {"columns": schema.names}

    def assert_single_bin(self, result, value, count, width):
        self.assertEqual(result["counts"], [count])
        self.assertEqual(len(result["bins"]), len(result["counts"]) + 1)
        lower, upper = result["bins"]
        self.assertGreaterEqual(lower, 0)
        self.assertLess(lower, upper)
        self.assertLessEqual(lower, value)
        self.assertLessEqual(value, upper)
        self.assertAlmostEqual(upper - lower, width)

    def test_constant_values_preserve_point_and_distinct_vessel_counts(self):
        self.write([self.point(i * 60, mmsi=413000001 + i // 2,
                               trajectory_id=1 + i // 2) for i in range(6)])
        self.assert_single_bin(analysis.speed_histogram(bin_width=1), 5, 6, 1)
        self.assert_single_bin(analysis.length_distribution(bin_width=10), 100, 3, 10)

    def test_constant_values_use_selected_or_default_bin_width(self):
        self.write([self.point()])
        self.assert_single_bin(analysis.speed_histogram(bin_width=2), 5, 1, 2)
        self.assert_single_bin(analysis.length_distribution(bin_width=20), 100, 1, 20)
        self.assert_single_bin(analysis.speed_histogram(), 5, 1, 1)
        self.assert_single_bin(analysis.length_distribution(), 100, 1, 10)

    def test_zero_speed_is_counted_without_negative_bin_boundaries(self):
        self.write([self.point(speed=0), self.point(60, speed=0)])
        result = analysis.speed_histogram(bin_width=1)
        self.assert_single_bin(result, 0, 2, 1)
        self.assertEqual(result["bins"], [0, 1])

    def test_missing_speed_and_invalid_lengths_are_not_counted(self):
        self.write([self.point(), self.point(60, speed=None, length=None),
                    self.point(120, length=0), self.point(180, length=501),
                    self.point(240, mmsi=413000002, trajectory_id=2),
                    self.point(300, length=float("inf")),
                    self.point(360, length=float("nan"))])
        self.assert_single_bin(analysis.speed_histogram(), 5, 6, 1)
        self.assert_single_bin(analysis.length_distribution(), 100, 2, 10)

    def test_empty_results_and_no_valid_lengths_stay_empty(self):
        self.write([self.point(length=0), self.point(60, length=None),
                    self.point(120, length=-1), self.point(180, length=501)])
        empty = {"bins": [], "counts": []}
        filters = QueryFilters(mmsi=999999999)
        self.assertEqual(analysis.speed_histogram(filters=filters), empty)
        self.assertEqual(analysis.length_distribution(filters=filters), empty)
        self.assertEqual(analysis.length_distribution(), empty)

    def test_counts_respect_time_vessel_and_region_filters(self):
        self.write([self.point(), self.point(60), self.point(120, lon=111),
                    self.point(180, mmsi=413000002, trajectory_id=2)])
        mode, geometry = normalize_spatial_filter("region", json.dumps(
            [[109.9, 19.9], [110.1, 19.9], [110.1, 20.1], [109.9, 20.1]]))
        region = QueryFilters(spatial_mode=mode, spatial_geometry=geometry)
        self.assert_single_bin(analysis.speed_histogram(filters=region), 5, 3, 1)
        self.assert_single_bin(analysis.length_distribution(filters=region), 100, 2, 10)
        restricted = QueryFilters(start_ts=self.start + 60, mmsi=413000001,
                                  spatial_mode=mode, spatial_geometry=geometry)
        self.assert_single_bin(analysis.speed_histogram(filters=restricted), 5, 1, 1)
        self.assert_single_bin(analysis.length_distribution(filters=restricted), 100, 1, 10)

    def test_section_counts_crossings_and_distinct_vessels(self):
        self.write([self.point(), self.point(60, lon=111), self.point(120),
                    self.point(mmsi=413000002, trajectory_id=2, lon=111),
                    self.point(60, mmsi=413000002, trajectory_id=2)])
        mode, geometry = normalize_spatial_filter("section", json.dumps(
            [[110.5, 19.5], [110.5, 20.5]]))
        section = QueryFilters(spatial_mode=mode, spatial_geometry=geometry)
        self.assert_single_bin(analysis.speed_histogram(filters=section), 5, 3, 1)
        self.assert_single_bin(analysis.length_distribution(filters=section), 100, 2, 10)

    def test_nonconstant_values_keep_existing_histogram_counts(self):
        self.write([self.point(i * 60, speed=float(i), length=100.0 + 100 * (i // 2),
                               mmsi=413000001 + i // 2, trajectory_id=1 + i // 2)
                    for i in range(6)])
        for result, count in ((analysis.speed_histogram(bin_width=1), 6),
                              (analysis.length_distribution(bin_width=10), 3)):
            self.assertEqual(sum(result["counts"]), count)
            self.assertGreater(len(result["counts"]), 1)
            self.assertEqual(len(result["bins"]), len(result["counts"]) + 1)

    def test_excel_exports_real_counts_and_one_hundred_percent(self):
        self.write([self.point(i * 60, mmsi=413000001 + i // 2,
                               trajectory_id=1 + i // 2) for i in range(6)])
        workbook = load_workbook(exporter.export_excel(), read_only=True, data_only=True)
        try:
            for name, count in (("航速分布", 6), ("船长分布", 3)):
                rows = list(workbook[name].iter_rows(min_row=3, values_only=True))
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0][2:], (count, 100))
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
