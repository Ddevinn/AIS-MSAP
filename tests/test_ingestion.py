from __future__ import annotations

import csv
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd
import pyarrow.parquet as pq

from backend.app.services import dataset
from backend.app.services.cleaning import run_cleaning
from backend.app.services.ingestion import (
    build_tabular_parquet, inspect_tabular, normalize_frame, parse_timestamps, resolve_fields,
)


class IngestionTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def write_csv(self, rows, headers=None, encoding="utf-8-sig"):
        path = self.root / "points.csv"
        headers = headers or ["MMSI", "Update_Time", "经度", "纬度", "SOG", "COG", "type", "船长", "船宽", "draft", "heading"]
        with path.open("w", encoding=encoding, newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(rows)
        return path

    def row(self, timestamp="2022-02-01 19:32:44", speed=9.7):
        return [413523230, timestamp, 110.12933, 20.21789, speed, 196.8, "-", 123, 20, 3.8, 206]

    def build(self, path):
        return build_tabular_parquet(path, self.root / "points.parquet", self.root / "manifest.json", "test", "test")

    def test_named_alias_mapping_ignores_order_and_preserves_heading(self):
        columns = [" heading ", "COG", "UpDate_Time", "MMSI", "LATITUDE", "LONGITUDE", "SOG", "船长", "beam", "draft", "TYPE"]
        mapping = resolve_fields(columns)
        self.assertEqual(mapping["timestamp"], "UpDate_Time")
        self.assertEqual(mapping["course"], "COG")
        self.assertEqual(mapping["heading"], " heading ")
        self.assertEqual(mapping["draught"], "draft")

    def test_ambiguous_alias_and_missing_required_columns_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "多个字段"):
            resolve_fields(["mmsi", "timestamp", "updatetime", "lon", "lat", "speed", "course"])
        with self.assertRaisesRegex(ValueError, "缺少必需字段"):
            resolve_fields(["mmsi", "lon", "lat"])
        with self.assertRaisesRegex(ValueError, "重复"):
            resolve_fields(["mmsi", "MMSI", "updatetime", "lon", "lat", "speed", "course"])

    def test_seconds_milliseconds_microseconds_nanoseconds_and_timezone(self):
        expected = 1643715164
        values = pd.Series([str(expected), str(expected * 1000), str(expected * 1000000), str(expected * 1000000000),
                            "2022-02-01 19:32:44", "2022-02-01T11:32:44Z", "2022-02-01T19:32:44+08:00",
                            datetime(2022, 2, 1, 19, 32, 44), "NULL", "not a date"])
        parsed = parse_timestamps(values, "Asia/Shanghai")
        self.assertEqual(parsed.tolist(), [expected] * 8 + [0, 0])
        utc = parse_timestamps(pd.Series(["2022-02-01 19:32:44"]), "UTC")
        self.assertEqual(utc.iloc[0], expected + 8 * 3600)

    def test_csv_utf8_and_gbk(self):
        for encoding in ["utf-8-sig", "gb18030"]:
            with self.subTest(encoding=encoding):
                path = self.write_csv([self.row()], encoding=encoding)
                preview = inspect_tabular(path)
                self.assertEqual(preview["mapping"]["timestamp"], "Update_Time")
                self.assertEqual(preview["preview"][0]["timestamp"], 1643715164)
                self.assertEqual(preview["preview"][0]["ship_type"], -1)

    def test_reordered_csv_becomes_typed_sorted_parquet(self):
        path = self.write_csv([self.row("2022-02-01 19:33:44"), self.row()])
        manifest = self.build(path)
        table = pq.read_table(self.root / "points.parquet")
        self.assertEqual(manifest["row_count"], 2)
        self.assertEqual(table["timestamp"].to_pylist(), [1643715164, 1643715224])
        self.assertAlmostEqual(table["course"][0].as_py(), 196.8, places=4)
        self.assertEqual(table["heading"][0].as_py(), 206)
        self.assertAlmostEqual(table["draught"][0].as_py(), 3.8, places=5)
        self.assertEqual(len(set(table["trajectory_id"].to_pylist())), 1)

    def test_invalid_values_remain_explicit_and_do_not_become_real_zero_measurements(self):
        columns = ["mmsi", "updatetime", "lon", "lat", "speed", "course", "length"]
        frame = pd.DataFrame([["413523230.5", "bad date", "NULL", "20", "bad speed", "90", "inf"]], columns=columns)
        normalized, issues = normalize_frame(frame, resolve_fields(columns), "Asia/Shanghai")
        self.assertEqual(normalized["mmsi"].iloc[0], 0)
        self.assertEqual(normalized["timestamp"].iloc[0], 0)
        self.assertTrue(pd.isna(normalized["speed"].iloc[0]))
        self.assertTrue(pd.isna(normalized["length"].iloc[0]))
        self.assertEqual(issues["mmsi"], 1)
        self.assertEqual(issues["timestamp"], 1)
        self.assertEqual(issues["length"], 1)

    def test_failed_rebuild_keeps_existing_parquet(self):
        path = self.write_csv([self.row()])
        self.build(path)
        before = (self.root / "points.parquet").read_bytes()
        self.write_csv([[1]], headers=["unrelated"])
        with self.assertRaises(ValueError):
            self.build(path)
        self.assertEqual((self.root / "points.parquet").read_bytes(), before)
        self.assertFalse((self.root / "points.parquet.import.tmp").exists())

    def test_registered_csv_build_and_cleaning_preserve_draught_and_heading(self):
        path = self.write_csv([self.row(), self.row("2022-02-01 19:33:44")])
        with patch.object(dataset, "DATASETS_DIR", self.root / "datasets"), \
             patch.object(dataset, "DATASET_REGISTRY_PATH", self.root / "datasets/registry.json"):
            record = dataset.register_dataset("csv test", path)
            manifest = dataset.build_parquet(dataset_id=record["id"])
            cleaned = run_cleaning(dataset_id=record["id"])
            self.assertEqual(manifest["row_count"], 2)
            self.assertEqual(cleaned["row_count"], 2)
            self.assertIn("draught", cleaned["columns"])
            self.assertIn("heading", cleaned["columns"])
            points = pq.read_table(dataset.cleaned_parquet_path(record["id"]))
            self.assertEqual(points["heading"].to_pylist(), [206, 206])
            self.assertAlmostEqual(points["draught"][0].as_py(), 3.8, places=5)

    def test_legacy_pickle_still_builds_and_cleans(self):
        path = self.root / "legacy.pkl"
        pd.to_pickle([np.asarray([[413523230, 1643715164, 110.12933, 20.21789, 196.8, 9.7, 70, 123, 20]], dtype=float)], path)
        with patch.object(dataset, "DATASETS_DIR", self.root / "datasets"), \
             patch.object(dataset, "DATASET_REGISTRY_PATH", self.root / "datasets/registry.json"):
            record = dataset.register_dataset("legacy test", path)
            self.assertEqual(dataset.build_parquet(dataset_id=record["id"])["row_count"], 1)
            self.assertEqual(run_cleaning(dataset_id=record["id"])["row_count"], 1)

    @unittest.skipUnless(os.environ.get("AIS_SAMPLE_XLSX"), "Set AIS_SAMPLE_XLSX to test the customer's workbook locally")
    def test_customer_excel_and_csv_produce_same_canonical_points(self):
        sample = Path(os.environ["AIS_SAMPLE_XLSX"])
        preview = inspect_tabular(sample)
        self.assertEqual(preview["mapping"]["timestamp"], "updatetime")
        self.assertEqual(preview["mapping"]["ship_type"], "type")
        self.assertEqual(preview["mapping"]["draught"], "draught")
        excel_manifest = self.build(sample)
        excel = pq.read_table(self.root / "points.parquet")
        self.assertEqual(excel_manifest["row_count"], 6116)
        from openpyxl import load_workbook
        workbook = load_workbook(sample, read_only=True, data_only=True)
        try:
            csv_path = self.root / "sample.csv"
            with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
                csv.writer(handle).writerows(workbook.worksheets[0].iter_rows(values_only=True))
        finally:
            workbook.close()
        self.build(csv_path)
        csv_points = pq.read_table(self.root / "points.parquet")
        self.assertTrue(excel.equals(csv_points))


if __name__ == "__main__":
    unittest.main()
