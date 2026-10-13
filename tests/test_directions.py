from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from backend.app.services import analysis, directions, exporter, query
from backend.app.services.query import QueryFilters, normalize_spatial_filter


class DirectionTests(unittest.TestCase):
    vessel = 413523230
    start = 1_700_000_000

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.parquet = self.root / "points.parquet"
        self.manifest = {}
        for target, name, value in (
            (analysis, "dataset_ready", lambda: True),
            (analysis, "read_active_manifest", lambda: self.manifest),
            (query, "active_parquet_path", lambda: self.parquet),
            (query, "read_active_manifest", lambda: self.manifest),
            (directions, "read_active_manifest", lambda: self.manifest),
        ):
            patcher = patch.object(target, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.addCleanup(self.temporary.cleanup)

    def point(self, course, heading, offset=0, **extra):
        row = {"mmsi": self.vessel, "timestamp": self.start + offset,
               "lon": 110.0, "lat": 20.0, "speed": 5.0, "course": course,
               "heading": heading, "ship_type": 70, "length": 100.0,
               "width": 20.0, "trajectory_id": 1}
        row.update(extra)
        return row

    def write_points(self, points, with_heading=True):
        schema = pa.schema([
            ("mmsi", pa.int64()), ("timestamp", pa.int64()),
            ("lon", pa.float64()), ("lat", pa.float64()), ("speed", pa.float64()),
            ("course", pa.float64()), ("ship_type", pa.int16()),
            ("length", pa.float64()), ("width", pa.float64()), ("trajectory_id", pa.int64()),
            *([("heading", pa.float64())] if with_heading else []),
        ])
        pq.write_table(pa.Table.from_pylist(points, schema=schema), self.parquet)
        self.manifest = {"columns": schema.names, "config": {"max_time_gap_minutes": 10}}

    def test_histograms_count_unavailable_and_invalid_values_and_use_circular_mean(self):
        courses = [359, 1, None, 360, 400, float("nan"), float("inf"), -5]
        headings = [10, 20, 511, None, -1, 361, float("nan"), float("inf")]
        self.write_points([self.point(c, h, index * 60) for index, (c, h) in enumerate(zip(courses, headings))])
        result = directions.direction_histograms()
        self.assertEqual(result["bins"], list(range(0, 361, 15)))
        self.assertEqual(result["point_count"], 8)
        for field in ("course", "heading"):
            self.assertEqual(result[field]["valid_count"], 2)
            self.assertEqual(result[field]["missing_count"], 2)
            self.assertEqual(result[field]["invalid_count"], 4)
            self.assertEqual(sum(result[field]["counts"]), 2)
            self.assertAlmostEqual(sum(result[field]["percentages"]), 100)
            self.assertEqual(result[field]["unusable_percentage"], 75)
        self.assertAlmostEqual(result["course"]["circular_mean"], 0)
        self.assertAlmostEqual(result["heading"]["circular_mean"], 15)
        self.assertEqual(result["course"]["counts"][0], 1)
        self.assertEqual(result["course"]["counts"][-1], 1)
        json.dumps(result, allow_nan=False)

    def test_opposite_directions_have_no_misleading_mean(self):
        self.write_points([self.point(0, 90), self.point(180, 270, 60)])
        result = directions.direction_histograms()
        self.assertIsNone(result["course"]["circular_mean"])
        self.assertIsNone(result["heading"]["circular_mean"])

    def test_histogram_and_curve_share_vessel_time_speed_and_spatial_filters(self):
        self.write_points([
            self.point(10, 20), self.point(30, 40, 60, lon=111),
            self.point(50, 60, 120, lon=111, mmsi=413523231, trajectory_id=2),
            self.point(70, 80, 180, lon=112), self.point(90, 100, 240),
        ])
        box = QueryFilters(mmsi=self.vessel, start_ts=self.start + 60,
                           end_ts=self.start + 180, speed_min=4, lon_min=110.5, lon_max=111.5)
        mode, geometry = normalize_spatial_filter("region", json.dumps([[110.5, 19.5], [111.5, 19.5], [111.5, 20.5], [110.5, 20.5]]))
        region = QueryFilters(mmsi=self.vessel, start_ts=self.start + 60,
                              end_ts=self.start + 180, spatial_mode=mode, spatial_geometry=geometry)
        mode, geometry = normalize_spatial_filter("section", json.dumps([[110.8, 19.5], [110.8, 20.5]]))
        section = QueryFilters(mmsi=self.vessel, end_ts=self.start + 180,
                               spatial_mode=mode, spatial_geometry=geometry)
        for filters in (box, region, section):
            with self.subTest(filters=filters):
                hist = directions.direction_histograms(filters=filters)
                curve = directions.direction_time_series(filters)
                self.assertEqual(hist["point_count"], 1)
                self.assertEqual(curve["total_points"], 1)
                self.assertEqual(curve["points"][0]["timestamp"], self.start + 60)
                self.assertEqual(curve["points"][0]["heading"], 40)
        empty = QueryFilters(mmsi=999999999)
        self.assertEqual(directions.direction_histograms(filters=empty)["point_count"], 0)
        self.assertEqual(directions.direction_time_series(empty)["points"], [])

    def test_legacy_dataset_without_heading_still_supports_course(self):
        self.write_points([self.point(90, None), self.point(100, None, 60)], with_heading=False)
        hist = directions.direction_histograms()
        curve = directions.direction_time_series(QueryFilters(mmsi=self.vessel))
        self.assertFalse(hist["heading"]["available"])
        self.assertEqual(hist["heading"]["missing_count"], 2)
        self.assertEqual(sum(hist["heading"]["counts"]), 0)
        self.assertEqual(hist["course"]["valid_count"], 2)
        self.assertFalse(curve["heading_available"])
        self.assertTrue(all(point["heading"] is None for point in curve["points"]))
        json.dumps(curve, allow_nan=False)

    def test_curve_breaks_each_field_at_wraps_missing_values_segments_and_time_gaps(self):
        self.write_points([
            self.point(359, 20), self.point(1, 30, 60), self.point(2, 350, 120),
            self.point(3, 10, 180), self.point(4, 511, 240), self.point(5, 20, 300),
            self.point(6, 21, 360, trajectory_id=2), self.point(7, 22, 2000, trajectory_id=2),
        ])
        curve = directions.direction_time_series(QueryFilters(mmsi=self.vessel))
        points = curve["points"]
        self.assertTrue(points[1]["course_break"])
        self.assertFalse(points[1]["heading_break"])
        self.assertFalse(points[2]["course_break"])
        self.assertTrue(points[2]["heading_break"])
        self.assertIsNone(points[4]["heading"])
        self.assertTrue(points[5]["heading_break"])
        for point in points[6:]:
            self.assertTrue(point["course_break"])
            self.assertTrue(point["heading_break"])

    def test_point_limit_preserves_endpoints_and_breaks_hidden_between_them(self):
        self.write_points([self.point(10, 30), self.point(359, 40, 60),
                           self.point(1, None, 120), self.point(20, 50, 180)])
        curve = directions.direction_time_series(QueryFilters(mmsi=self.vessel), max_points=2)
        self.assertEqual(curve["total_points"], 4)
        self.assertEqual(curve["displayed_points"], 2)
        self.assertTrue(curve["sampled"])
        self.assertEqual([point["timestamp"] for point in curve["points"]], [self.start, self.start + 180])
        self.assertTrue(curve["points"][-1]["course_break"])
        self.assertTrue(curve["points"][-1]["heading_break"])

    def test_excel_adds_three_sheets_with_filtered_counts_and_selected_bins(self):
        self.write_points([self.point(10, 20), self.point(80, 511, 60),
                           self.point(180, 270, 120, mmsi=413523231, trajectory_id=2)])
        with patch.object(exporter, "EXPORT_DIR", self.root / "exports"):
            output = exporter.export_excel(QueryFilters(mmsi=self.vessel), direction_bin_width=45)
        workbook = load_workbook(output, read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(len(workbook.sheetnames), 9)
        overview = workbook["航向统计"]
        self.assertEqual(overview["B4"].value, 2)
        self.assertEqual(overview["C5"].value, 1)
        self.assertEqual(overview["C6"].value, 1)
        for title, expected in (("对地航向分布", 2), ("船首向分布", 1)):
            sheet = workbook[title]
            self.assertEqual(sheet.max_row, 10)
            self.assertEqual(sheet["B3"].value, 45)
            self.assertEqual(sum(sheet.cell(row, 3).value for row in range(3, 11)), expected)

    def test_bad_bins_and_missing_vessel_are_rejected(self):
        self.write_points([self.point(10, 20)])
        with self.assertRaises(ValueError):
            directions.direction_histograms(bin_width=7)
        with self.assertRaises(ValueError):
            directions.direction_time_series(QueryFilters())


if __name__ == "__main__":
    unittest.main()
