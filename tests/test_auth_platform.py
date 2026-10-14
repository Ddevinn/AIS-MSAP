from __future__ import annotations

import io
import json
import tempfile
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient

from backend.app import main, storage
from backend.app.services import dataset, exporter, raster_tracks


class AccountPlatformTests(unittest.TestCase):
    """Exercise real HTTP authorization, files and query threads in a temporary store."""
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        for module, field, value in (
            (storage, "DB_PATH", self.root / "private" / "test.sqlite3"),
            (storage, "DATASET_REGISTRY_PATH", self.root / "registry.json"),
            (dataset, "DATASETS_DIR", self.root / "datasets"),
            (main, "DATA_DIR", self.root),
            (exporter, "EXPORT_DIR", self.root / "exports"),
            (raster_tracks, "CACHE_DIR_PATH", self.root / "cache"),
        ):
            patcher = patch.object(module, field, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.addCleanup(self.temp.cleanup)
        storage.initialize()
        self.admin = self.client()
        response = self.admin.post("/api/auth/setup", json={"username": "admin", "password": "AdminPassword123!", "display_name": "管理员"})
        self.assertEqual(response.status_code, 200, response.text)
        self.login(self.admin, "admin", "AdminPassword123!")

    def client(self):
        client = TestClient(main.app, base_url="http://127.0.0.1", client=("127.0.0.1", 12345))
        self.addCleanup(client.close)
        return client

    def login(self, client, username, password):
        response = client.post("/api/auth/login", json={"username": username, "password": password})
        self.assertEqual(response.status_code, 200, response.text)
        client.headers["X-CSRF-Token"] = response.json()["csrf_token"]
        return response.json()

    def user(self, username="alice"):
        response = self.admin.post("/api/admin/users", json={"username": username, "password": "TemporaryPass123!", "display_name": username})
        self.assertEqual(response.status_code, 200, response.text)
        client = self.client()
        self.login(client, username, "TemporaryPass123!")
        response = client.post("/api/auth/password", json={"current_password": "TemporaryPass123!", "new_password": "ChangedPassword123!"})
        self.assertEqual(response.status_code, 200, response.text)
        self.login(client, username, "ChangedPassword123!")
        return client, self.admin.get("/api/admin/users").json()["users"][-1]["id"]

    def upload(self, client, count=3, name="船舶数据"):
        contents = "mmsi,updatetime,lon,lat,speed,course,heading,ship_type,length,width\n"
        contents += "".join(f"413523230,2024-01-01 12:0{i}:00,{110 + i * 0.001},20,5,90,91,70,100,20\n" for i in range(count))
        response = client.post("/api/datasets/upload", data={"name": name}, files={"file": ("points.csv", contents.encode(), "text/csv")})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["dataset"]["id"]

    def build(self, client):
        response = client.post("/api/build")
        self.assertEqual(response.status_code, 200, response.text)
        end = time.monotonic() + 10
        while time.monotonic() < end:
            state = client.get("/api/status").json()
            if state["build_job"]["status"] != "running":
                self.assertEqual(state["build_job"]["status"], "done", state)
                return state
            time.sleep(0.03)
        self.fail("Temporary dataset build did not finish")

    def test_pages_and_every_data_entry_require_login(self):
        anonymous = self.client()
        for path in ("/api/status", "/api/datasets", "/api/summary", "/api/density-grid",
                     "/api/tiles/tracks/8/206/113.png", "/api/export/excel", "/api/admin/users", "/api/cleaning/status"):
            self.assertEqual(anonymous.get(path).status_code, 401, path)
        for path in ("/", "/admin", "/static/index.html", "/static/admin.html", "/static//admin.html", "/static/ADMIN.HTML"):
            response = anonymous.get(path, follow_redirects=False)
            self.assertEqual(response.status_code, 303, path)
        self.assertEqual(anonymous.get("/login").status_code, 200)

    def test_empty_login_and_history_survive_relogin_without_auto_selection(self):
        alice, _ = self.user()
        initial = alice.get("/api/status").json()
        self.assertIsNone(initial["active_dataset_id"])
        self.assertEqual(initial["datasets"], [])
        self.assertFalse(initial["dataset_ready"])
        self.assertEqual(alice.post("/api/build").status_code, 409)
        dataset_id = self.upload(alice)
        self.build(alice)
        self.assertEqual(alice.get("/api/summary").json()["point_count"], 3)
        # Refresh keeps the current login's selection.
        self.assertEqual(alice.get("/api/status").json()["active_dataset_id"], dataset_id)
        self.assertEqual(alice.post("/api/auth/logout").status_code, 200)
        self.login(alice, "alice", "ChangedPassword123!")
        state = alice.get("/api/status").json()
        self.assertIsNone(state["active_dataset_id"])
        self.assertEqual(len(state["datasets"]), 1)
        self.assertEqual(state["datasets"][0]["filename"], "points.csv")
        response = alice.post(f"/api/datasets/select/{dataset_id}")
        self.assertEqual(response.status_code, 200)
        parquet = Path(alice.get("/api/status").json()["active_parquet"])
        modified = parquet.stat().st_mtime_ns
        self.assertFalse(alice.post("/api/build").json()["accepted"])
        self.assertEqual(parquet.stat().st_mtime_ns, modified)
        self.assertEqual(alice.post("/api/datasets/clear").status_code, 200)
        self.assertEqual(alice.delete(f"/api/datasets/{dataset_id}").status_code, 200)
        self.assertEqual(alice.get("/api/datasets").json()["datasets"], [])
        self.assertTrue(parquet.exists())

    def test_owned_datasets_and_admin_are_inaccessible_to_other_accounts(self):
        alice, _ = self.user("alice")
        bob, _ = self.user("bob")
        alice_id = self.upload(alice)
        self.build(alice)
        self.assertEqual(bob.get("/api/datasets").json()["datasets"], [])
        for method, path in (("post", f"/api/datasets/select/{alice_id}"), ("delete", f"/api/datasets/{alice_id}"),
                             ("get", f"/api/summary?dataset_id={alice_id}"),
                             ("get", f"/api/tiles/tracks/8/206/113.png?dataset_id={alice_id}"),
                             ("get", f"/api/export/excel?dataset_id={alice_id}")):
            self.assertEqual(getattr(bob, method)(path).status_code, 404, path)
        for path in ("/admin", "/static//admin.html", "/static/ADMIN.HTML", "/api/admin/users", "/api/admin/sessions", "/api/admin/logs", "/openapi.json"):
            self.assertEqual(alice.get(path).status_code, 403, path)
        self.assertEqual(alice.post("/api/datasets/register", json={"source_path": str(self.root)}).status_code, 403)
        self.assertEqual(alice.post("/api/datasets/preview", json={"source_path": str(self.root)}).status_code, 403)

    def test_parallel_queries_rendering_jobs_and_export_stay_isolated(self):
        alice, _ = self.user("alice")
        bob, _ = self.user("bob")
        alice_id, bob_id = self.upload(alice, 2, "同名数据"), self.upload(bob, 5, "同名数据")
        self.assertNotEqual(alice_id, bob_id)
        self.build(alice)
        self.build(bob)
        with ThreadPoolExecutor(max_workers=4) as pool:
            futures = [pool.submit(client.get, "/api/summary") for client in (alice, bob, alice, bob)]
            self.assertEqual([future.result().json()["point_count"] for future in futures], [2, 5, 2, 5])
        # Exercise real run_in_executor rendering: lost context would read the legacy global registry.
        for client, dataset_id in ((alice, alice_id), (bob, bob_id)):
            tile = client.get(f"/api/tiles/tracks/8/206/113.png?dataset_id={dataset_id}")
            self.assertEqual(tile.status_code, 200, tile.text[:100] if tile.status_code != 200 else "")
            self.assertTrue(tile.content.startswith(b"\x89PNG"))
            self.assertEqual(tile.headers["Cache-Control"], "no-store")
            response = client.get("/api/export/excel")
            self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else "")
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(response.content), read_only=True)
            values = dict(workbook["总览"].iter_rows(min_row=3, max_col=2, values_only=True))
            self.assertEqual(values["数据点/截面穿越次数"], 2 if client is alice else 5)
            workbook.close()
        self.assertTrue((self.root / "exports" / "user_2").exists())
        self.assertTrue((self.root / "exports" / "user_3").exists())

    def test_two_sessions_of_same_user_have_independent_active_data(self):
        alice, _ = self.user()
        first = self.upload(alice, 2)
        self.build(alice)
        second = self.upload(alice, 4)
        self.build(alice)
        other = self.client()
        self.login(other, "alice", "ChangedPassword123!")
        self.assertIsNone(other.get("/api/status").json()["active_dataset_id"])
        other.post(f"/api/datasets/select/{first}")
        self.assertEqual(alice.get("/api/status").json()["active_dataset_id"], second)
        self.assertEqual(other.get("/api/summary").json()["point_count"], 2)
        self.assertEqual(alice.get("/api/summary").json()["point_count"], 4)

    def test_cleaning_uses_owner_job_and_never_changes_other_users_data(self):
        alice, _ = self.user("alice")
        bob, _ = self.user("bob")
        self.upload(alice); self.build(alice)
        self.upload(bob, 4); self.build(bob)
        response = alice.post("/api/cleaning/run", json={})
        self.assertEqual(response.status_code, 200, response.text)
        end = time.monotonic() + 10
        while time.monotonic() < end:
            state = alice.get("/api/status").json()
            if state["clean_job"]["status"] != "running":
                break
            time.sleep(0.03)
        self.assertEqual(state["clean_job"]["status"], "done", state)
        self.assertTrue(state["cleaned_dataset_ready"])
        other = bob.get("/api/status").json()
        self.assertFalse(other["cleaned_dataset_ready"])
        self.assertEqual(other["clean_job"]["status"], "idle")
        self.assertEqual(bob.get("/api/summary").json()["point_count"], 4)

    def test_temp_password_restricts_features_until_change_and_invalidates_sessions(self):
        response = self.admin.post("/api/admin/users", json={"username": "alice", "password": "TemporaryPass123!"})
        self.assertEqual(response.status_code, 200)
        alice = self.client()
        self.login(alice, "alice", "TemporaryPass123!")
        response = alice.get("/api/status")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json()["code"], "PASSWORD_CHANGE_REQUIRED")
        response = alice.get("/", follow_redirects=False)
        self.assertEqual(response.headers["location"], "/account/password")
        self.assertEqual(alice.post("/api/auth/password", json={"current_password": "wrong", "new_password": "NewPassword123!"}).status_code, 400)

    def test_password_reset_disable_and_revoke_immediately_end_sessions(self):
        alice, user_id = self.user()
        reset = self.admin.post(f"/api/admin/users/{user_id}/password", json={"password": "ResetPassword123!", "must_change_password": False})
        self.assertEqual(reset.status_code, 200)
        self.assertEqual(alice.get("/api/auth/me").status_code, 401)
        self.assertEqual(alice.post("/api/auth/login", json={"username": "alice", "password": "ChangedPassword123!"}).status_code, 401)
        self.login(alice, "alice", "ResetPassword123!")
        sessions = self.admin.get(f"/api/admin/sessions?user_id={user_id}").json()["sessions"]
        self.assertEqual(self.admin.delete(f"/api/admin/sessions/{sessions[0]['id']}").status_code, 200)
        self.assertEqual(alice.get("/api/status").status_code, 401)
        self.login(alice, "alice", "ResetPassword123!")
        self.assertEqual(self.admin.patch(f"/api/admin/users/{user_id}", json={"display_name": "Alice", "role": "user", "enabled": False}).status_code, 200)
        self.assertEqual(alice.get("/api/status").status_code, 401)
        self.assertEqual(alice.post("/api/auth/login", json={"username": "alice", "password": "ResetPassword123!"}).status_code, 401)

    def test_csrf_origin_and_expiry_are_enforced(self):
        alice, _ = self.user()
        saved = alice.headers.pop("X-CSRF-Token")
        self.assertEqual(alice.post("/api/auth/logout").status_code, 403)
        alice.headers["X-CSRF-Token"] = saved
        self.assertEqual(alice.post("/api/auth/logout", headers={"Origin": "https://evil.example"}).status_code, 403)
        self.assertEqual(alice.post("/api/auth/logout", headers={"Origin": "http://127.0.0.1:invalid"}).status_code, 403)
        with storage.connection() as con:
            con.execute("UPDATE sessions SET expires_at=0 WHERE user_id=2")
        self.assertEqual(alice.get("/api/status").status_code, 401)

    def test_startup_recovers_jobs_and_password_reset_wins_login_race(self):
        alice, user_id = self.user()
        storage.save_job(user_id, "build", {"status": "running", "progress": {}})
        with TestClient(main.app, base_url="http://127.0.0.1", client=("127.0.0.1", 123)) as started:
            self.assertEqual(started.get("/api/auth/setup-status").status_code, 200)
            self.assertEqual(storage.load_job(user_id, "build")["status"], "error")
        from backend.app.auth import password_hash
        real_verify = password_hash.verify
        replacement = password_hash.hash("ReplacementPass123!")
        def reset_during_verify(password, hashed):
            valid = real_verify(password, hashed)
            with storage.connection() as con:
                con.execute("UPDATE users SET password_hash=? WHERE id=?", (replacement, user_id))
                con.execute("DELETE FROM sessions WHERE user_id=?", (user_id,))
            return valid
        with patch.object(password_hash, "verify", reset_during_verify):
            response = alice.post("/api/auth/login", json={"username": "alice", "password": "ChangedPassword123!"})
        self.assertEqual(response.status_code, 401)
        self.assertEqual(alice.get("/api/status").status_code, 401)

    def test_admin_records_never_expose_passwords_or_session_secrets(self):
        alice, user_id = self.user()
        alice.get("/")
        response = self.admin.get("/api/admin/users")
        for user in response.json()["users"]:
            self.assertNotIn("password_hash", user)
            self.assertNotIn("password", user)
        with storage.connection() as con:
            hashed = con.execute("SELECT password_hash FROM users WHERE id=?", (user_id,)).fetchone()[0]
        self.assertTrue(hashed.startswith("$argon2"))
        self.assertNotIn("ChangedPassword123!", hashed)
        for kind in ("login", "access", "audit"):
            result = self.admin.get(f"/api/admin/logs?kind={kind}").json()
            self.assertGreater(result["total"], 0)
            self.assertNotIn("Password123", json.dumps(result))
        sessions = self.admin.get("/api/admin/sessions").json()["sessions"]
        self.assertTrue(sessions)
        for session in sessions:
            self.assertNotIn("csrf_token", session)
            self.assertNotIn("token", session)
        self.assertEqual(self.admin.patch("/api/admin/users/1", json={"display_name": "管理员", "role": "user", "enabled": True}).status_code, 400)
        self.assertEqual(self.admin.patch("/api/admin/users/1", json={"display_name": "管理员", "role": "superadmin", "enabled": False}).status_code, 400)

    def test_initialization_is_one_time_and_local_and_legacy_migration_is_read_only(self):
        self.assertEqual(self.admin.post("/api/auth/setup", json={"username": "again", "password": "AnotherPassword123!"}).status_code, 409)
        second = self.root / "migration.sqlite3"
        original = {"active_dataset_id": "legacy", "datasets": {"legacy": {"id": "legacy", "name": "旧数据", "source_path": str(self.root / "old.csv"),
            "source_parquet": str(self.root / "old.parquet"), "source_manifest": str(self.root / "manifest.json"),
            "cleaned_parquet": str(self.root / "cleaned.parquet"), "cleaning_manifest": str(self.root / "cleaning.json")}}}
        storage.DATASET_REGISTRY_PATH.write_text(json.dumps(original), encoding="utf-8")
        before = storage.DATASET_REGISTRY_PATH.read_bytes()
        with patch.object(storage, "DB_PATH", second):
            storage.initialize()
            remote = TestClient(main.app, base_url="http://127.0.0.1", client=("192.168.1.10", 123))
            self.addCleanup(remote.close)
            self.assertEqual(remote.post("/api/auth/setup", json={"username": "root", "password": "AdminPassword123!"}).status_code, 403)
            local = self.client()
            self.assertEqual(local.post("/api/auth/setup", json={"username": "root", "password": "AdminPassword123!"}).json()["migrated_datasets"], 1)
            self.login(local, "root", "AdminPassword123!")
            result = local.get("/api/datasets").json()
            self.assertIsNone(result["active_dataset_id"])
            self.assertEqual(result["datasets"][0]["name"], "旧数据")
        self.assertEqual(storage.DATASET_REGISTRY_PATH.read_bytes(), before)

    def test_restart_marks_interrupted_jobs_and_admin_login_rejects_regular_user(self):
        alice, user_id = self.user()
        response = alice.post("/api/auth/login", headers={"X-Login-Target": "admin"}, json={"username": "alice", "password": "ChangedPassword123!"})
        self.assertEqual(response.status_code, 401)
        storage.save_job(user_id, "build", {"status": "running", "dataset_id": "example", "progress": {"percent": 12}})
        storage.recover_jobs()
        self.assertEqual(storage.load_job(user_id, "build")["status"], "error")
        self.assertEqual(self.admin.get("/api/status").json()["build_job"]["status"], "idle")


if __name__ == "__main__":
    unittest.main()
